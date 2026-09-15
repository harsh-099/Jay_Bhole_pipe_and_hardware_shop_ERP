from django.shortcuts import render, redirect
from django.core.paginator import Paginator
from .models import Product
from .forms import ProductForm
from django.http import HttpResponse
import openpyxl
from django.contrib import messages
from openpyxl import load_workbook
from accounts.decorators import allowed_roles

@allowed_roles(["Admin","Manager","Accountant",])
def product_list(request):

    search = request.GET.get("search")

    products = Product.objects.all().order_by("name")

    if search:

        products = products.filter(
            name__icontains=search
        )

    paginator = Paginator(
        products,
        10
    )

    page_number = request.GET.get("page")

    page_obj = paginator.get_page(page_number)

    context = {

        "products": page_obj,

        "search": search,

        "total_products": Product.objects.count()

    }

    return render(
        request,
        "products/product_list.html",
        context
    )


@allowed_roles(["Admin","Manager","Accountant",])
def add_product(request):

    if request.method == "POST":

        form = ProductForm(request.POST)

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Product Added Successfully."
            )

            return redirect("product_list")

        else:

            messages.error(
                request,
                "Product already exists."
            )

    else:

        form = ProductForm()

    return render(
        request,
        "products/add_product.html",
        {
            "form": form
        }
    )

@allowed_roles(["Admin","Manager","Accountant",])
def edit_product(request, id):

    product = Product.objects.get(id=id)

    if request.method == "POST":

        form = ProductForm(
            request.POST,
            instance=product
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Product Updated Successfully."
            )

            return redirect("product_list")

        else:

            messages.error(
                request,
                "Another product with this name already exists."
            )

    else:

        form = ProductForm(
            instance=product
        )

    return render(
        request,
        "products/edit_product.html",
        {
            "form": form
        }
    )

@allowed_roles(["Admin","Manager","Accountant",])
def delete_product(request, id):

    product = Product.objects.get(id=id)

    product.delete()

    return redirect("product_list")

@allowed_roles(["Admin","Manager","Accountant",])
def export_products_excel(request):

    wb = openpyxl.Workbook()

    ws = wb.active

    ws.title = "Products"

    ws.append([
        "ID",
        "Product",
        "Category",
        "Purchase Price",
        "Selling Price",
        "Stock",
        "Unit"
    ])

    products = Product.objects.all()

    for p in products:

        ws.append([
            p.id,
            p.name,
            p.category,
            p.purchase_price,
            p.selling_price,
            p.quantity,
            p.unit
        ])

    response = HttpResponse(

        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    )

    response["Content-Disposition"] = 'attachment; filename="Products.xlsx"'

    wb.save(response)

    return response

from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from django.contrib import messages
from django.shortcuts import redirect, render

from accounts.decorators import allowed_roles
from .models import Product


def to_decimal(value):
    """
    Excel मधील value Decimal मध्ये convert करतो.
    Result फक्त 2 decimal places मध्ये ठेवतो.
    """

    if value is None or value == "":
        return Decimal("0.00")

    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0.00")


@allowed_roles(["Admin", "Manager", "Accountant"])
def import_products_excel(request):

    if request.method != "POST":
        return render(
            request,
            "products/import_products.html"
        )

    excel_file = request.FILES.get("excel_file")

    if not excel_file:
        messages.error(request, "Please Select Excel File.")
        return redirect("product_list")

    try:
        wb = load_workbook(excel_file, data_only=True)
        sheet = wb.active

    except Exception as e:
        messages.error(
            request,
            f"Excel file read error: {e}"
        )
        return redirect("product_list")

    imported = 0
    updated = 0
    skipped = 0

    for row_no, row in enumerate(
        sheet.iter_rows(min_row=2, values_only=True),
        start=2
    ):

        try:

            # Product name
            if len(row) < 7 or row[1] is None:
                skipped += 1
                continue

            name = str(row[1]).strip()

            # Category
            category = (
                str(row[2]).strip()
                if row[2] is not None
                else ""
            )

            # Purchase Price
            purchase_price = to_decimal(row[3])

            # Selling Price
            selling_price = to_decimal(row[4])

            # Quantity
            try:
                quantity = int(float(row[5] or 0))
            except (ValueError, TypeError):
                quantity = 0

            # Unit
            unit = (
                str(row[6]).strip()
                if row[6] is not None
                else ""
            )

            # Find existing product
            product = Product.objects.filter(
                name__iexact=name
            ).first()

            if product:

                product.category = category
                product.purchase_price = purchase_price
                product.selling_price = selling_price
                product.quantity = quantity
                product.unit = unit

                product.save()

                updated += 1

            else:

                Product.objects.create(
                    name=name,
                    category=category,
                    purchase_price=purchase_price,
                    selling_price=selling_price,
                    quantity=quantity,
                    unit=unit
                )

                imported += 1

            print(
                f"Row {row_no}:",
                name,
                category,
                purchase_price,
                selling_price,
                quantity,
                unit
            )

        except Exception as e:

            skipped += 1

            print(
                f"Row {row_no} Error -> {e}"
            )

    messages.success(
        request,
        f"Imported : {imported} | "
        f"Updated : {updated} | "
        f"Skipped : {skipped}"
    )

    return redirect("product_list")

@allowed_roles(["Admin", "Manager", "Accountant"])
def low_stock_products(request):

    products = Product.objects.filter(
        quantity__lte=10
    ).order_by("quantity", "name")

    return render(
        request,
        "products/low_stock.html",
        {
            "products": products
        }
    )